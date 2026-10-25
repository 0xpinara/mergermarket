"""Parse uploaded CSV/XLSX exports, detect their entity type and suggest a
column mapping. Nothing here evaluates cell content; values are kept as text."""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from rapidfuzz import fuzz

from . import config
from . import normalize as nz
from .schema import ENTITIES, ENTITY_ORDER, EntitySpec

ALLOWED_EXT = {".csv", ".tsv", ".txt", ".xlsx"}


class UploadError(ValueError):
    """A user-facing problem with an uploaded file."""


@dataclass
class Table:
    file_name: str
    sheet: str | None
    headers: list[str]
    rows: list[list[str]]
    row_numbers: list[int]
    warnings: list[str] = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"{self.file_name} › {self.sheet}" if self.sheet else self.file_name


def sanitize_filename(name: str | None) -> str:
    base = (name or "upload").replace("\\", "/").split("/")[-1]
    base = unicodedata.normalize("NFKC", base)
    base = re.sub(r"[^A-Za-z0-9._ \-()]+", "_", base).strip(" .") or "upload"
    return base[:120]


def _ext(name: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", name)
    return m.group(1).lower() if m else ""


def parse_upload(file_name: str, data: bytes) -> list[Table]:
    name = sanitize_filename(file_name)
    ext = _ext(name)
    if ext == ".xls":
        raise UploadError(f"{name}: legacy .xls workbooks aren't supported. Save as .xlsx or .csv and upload again.")
    if ext not in ALLOWED_EXT:
        raise UploadError(f"{name}: unsupported file type. Upload .csv or .xlsx exports.")
    if not data:
        raise UploadError(f"{name}: the file is empty.")
    if len(data) > config.MAX_UPLOAD_BYTES:
        mb = config.MAX_UPLOAD_BYTES // (1024 * 1024)
        raise UploadError(f"{name}: file is larger than the {mb} MB limit for this demo.")
    if ext == ".xlsx":
        return _parse_xlsx(name, data)
    return [_parse_csv(name, data)]


def _decode(name: str, data: bytes) -> str:
    if b"\x00" in data[:2048]:
        raise UploadError(f"{name}: this doesn't look like a text CSV (binary content found).")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _looks_like_value(h: str) -> bool:
    h = h.strip()
    if not h:
        return False
    if re.fullmatch(r"[-+$(]?\d[\d,]*(\.\d+)?\)?", h):
        return True
    return nz.parse_date(h) is not None or nz.email(h) is not None


def _finish(name: str, sheet: str | None, raw_rows: list[tuple[int, list[str]]]) -> Table:
    raw_rows = [(n, r) for n, r in raw_rows if any(c.strip() for c in r)]
    if not raw_rows:
        raise UploadError(f"{_label(name, sheet)}: no rows found.")
    header_line, header = raw_rows[0]
    header = [h.strip() for h in header]
    while header and not header[-1]:
        header.pop()
    if not header or sum(1 for h in header if h) == 0:
        raise UploadError(f"{_label(name, sheet)}: the first row has no column headers.")
    filled = [h for h in header if h]
    if sum(_looks_like_value(h) for h in filled) >= max(1, len(filled) // 2):
        raise UploadError(f"{_label(name, sheet)}: the first row looks like data, not column headers. "
                          "Add a header row (e.g. lease_id, unit_id, monthly_rent) and upload again.")
    if len(header) > config.MAX_COLUMNS:
        raise UploadError(f"{_label(name, sheet)}: more than {config.MAX_COLUMNS} columns.")
    seen: dict[str, int] = {}
    headers = []
    for i, h in enumerate(header):
        h = h or f"column_{i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        headers.append(h)
    data_rows = raw_rows[1:]
    if not data_rows:
        raise UploadError(f"{_label(name, sheet)}: has a header row but no data rows.")
    if len(data_rows) > config.MAX_ROWS_PER_FILE:
        raise UploadError(f"{_label(name, sheet)}: {len(data_rows):,} rows exceeds the "
                          f"{config.MAX_ROWS_PER_FILE:,}-row limit per file for this demo.")
    warnings = []
    ragged = 0
    rows, numbers = [], []
    width = len(headers)
    for n, r in data_rows:
        if len(r) != width:
            if len(r) > width and any(c.strip() for c in r[width:]):
                ragged += 1
            r = (r + [""] * width)[:width]
        rows.append([c.strip() for c in r])
        numbers.append(n)
    if ragged:
        warnings.append(f"{ragged} row(s) had more cells than headers; extra cells were ignored.")
    return Table(name, sheet, headers, rows, numbers, warnings)


def _label(name: str, sheet: str | None) -> str:
    return f"{name} › {sheet}" if sheet else name


def _parse_csv(name: str, data: bytes) -> Table:
    text = _decode(name, data)
    sample = text[:8192]
    delim = ","
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        if name.lower().endswith(".tsv"):
            delim = "\t"
    try:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
        raw = [(i + 1, row) for i, row in enumerate(reader)]
    except csv.Error as exc:
        raise UploadError(f"{name}: couldn't read this CSV ({exc}). Check quoting and delimiters.") from exc
    return _finish(name, None, raw)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)


def _parse_xlsx(name: str, data: bytes) -> list[Table]:
    if not data.startswith(b"PK"):
        raise UploadError(f"{name}: not a valid .xlsx workbook.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(i.file_size for i in zf.infolist())
            if total > config.MAX_XLSX_UNCOMPRESSED:
                raise UploadError(f"{name}: workbook expands to more than "
                                  f"{config.MAX_XLSX_UNCOMPRESSED // (1024 * 1024)} MB; too large for this demo.")
    except zipfile.BadZipFile as exc:
        raise UploadError(f"{name}: not a valid .xlsx workbook.") from exc
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except UploadError:
        raise
    except Exception as exc:  # openpyxl raises a variety of types on corrupt files
        raise UploadError(f"{name}: couldn't open this workbook. Re-save it as .xlsx or export CSV.") from exc
    tables, errors = [], []
    try:
        for ws in wb.worksheets:
            raw = []
            # Cap columns read per row: a single stray cell far to the right would otherwise make openpyxl
            # pad every row out to it (a few-KB workbook can expand to gigabytes).
            for i, row in enumerate(ws.iter_rows(values_only=True, max_col=config.MAX_COLUMNS + 1)):
                if i > config.MAX_ROWS_PER_FILE + 1:
                    raise UploadError(f"{name} › {ws.title}: exceeds the {config.MAX_ROWS_PER_FILE:,}-row limit.")
                if len(row) > config.MAX_COLUMNS and row[config.MAX_COLUMNS] not in (None, ""):
                    raise UploadError(f"{name} › {ws.title}: more than {config.MAX_COLUMNS} columns.")
                raw.append((i + 1, [_cell(v) for v in row[: config.MAX_COLUMNS]]))
            if not any(any(c.strip() for c in r) for _, r in raw):
                continue
            try:
                tables.append(_finish(name, ws.title if len(wb.worksheets) > 1 else None, raw))
            except UploadError as exc:
                errors.append(str(exc))
    finally:
        wb.close()
    if not tables:
        raise UploadError(errors[0] if errors else f"{name}: the workbook has no data.")
    if errors:
        tables[0].warnings.extend(f"Skipped sheet — {e}" for e in errors)
    return tables


# --------------------------------------------------------------------------- mapping

_NOISE = {"usd", "dollars", "dollar", "amt", "the", "value", "field", "col", "column", "of", "in", "s"}


def norm_header(h: str) -> str:
    h = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", h)
    h = "".join(c for c in unicodedata.normalize("NFKD", h) if not unicodedata.combining(c)).lower()
    h = re.sub(r"[^a-z0-9]+", " ", h)
    return re.sub(r"\s+", " ", h).strip()


def _strip_noise(h: str, entity: EntitySpec | None = None) -> str:
    drop = set(_NOISE)
    if entity:
        drop |= {entity.singular, entity.key, entity.key.rstrip("s")}
    toks = [t for t in h.split() if t not in drop]
    return " ".join(toks)


def _stem(t: str) -> str:
    for suf in ("ing", "ed", "es", "s"):
        if len(t) > len(suf) + 2 and t.endswith(suf):
            t = t[: -len(suf)]
            break
    return t[:-1] if len(t) > 3 and t.endswith("e") else t


def _stems(h: str) -> set[str]:
    return {_stem(t) for t in h.split() if t not in {"of", "the", "a", "on", "by"}}


def suggest_mapping(entity_key: str, headers: list[str]) -> dict[str, dict]:
    """Return {field_key: {"column", "confidence", "method"}} for the best guesses."""
    entity = ENTITIES[entity_key]
    normed = {h: norm_header(h) for h in headers}
    cands: list[tuple[float, int, int, str, str, str]] = []
    for f_idx, f in enumerate(entity.fields):
        aliases = [norm_header(a) for a in f.aliases]
        for h, nh in normed.items():
            best: tuple[float, str] | None = None
            if nh in aliases:
                best = (1.0, "exact alias")
            else:
                sh = _strip_noise(nh)
                if sh and sh in aliases:
                    best = (0.9, "alias (ignoring units/noise words)")
                else:
                    sh2 = _strip_noise(nh, entity)
                    if sh2 and sh2 in [_strip_noise(a, entity) for a in aliases]:
                        best = (0.8, f"alias (ignoring '{entity.singular}' prefix)")
                    else:
                        # every word of a multi-word alias appears in the header ("Resident Moved In" ⊇ "move in")
                        hs = _stems(nh)
                        hit = max((a for a in aliases if len(a.split()) >= 2 and _stems(a) <= hs),
                                  key=lambda a: len(a.split()), default=None)
                        if hit:
                            best = (0.72, f"alias words in header ('{hit}')")
                        else:
                            score = max((fuzz.ratio(nh, a) for a in aliases), default=0)
                            if score >= 88:
                                best = (round(score / 100 * 0.7, 2), "similar spelling")
            if best:
                alias_rank = next((i for i, a in enumerate(aliases) if a == nh), len(aliases))
                cands.append((best[0], 0 if f.required else 1, alias_rank, f.key, h, best[1]))
    cands.sort(key=lambda c: (-c[0], c[1], c[2], entity.fields.index(entity.field(c[3]))))
    mapping: dict[str, dict] = {}
    used_cols: set[str] = set()
    for conf, _req, _rank, fkey, col, method in cands:
        if fkey in mapping or col in used_cols:
            continue
        mapping[fkey] = {"column": col, "confidence": conf, "method": method}
        used_cols.add(col)
    return mapping


def detect_entity(file_name: str, headers: list[str]) -> tuple[str, float, dict[str, float]]:
    fname = norm_header(re.sub(r"\.[A-Za-z0-9]+$", "", file_name)).replace(" ", "_")
    scores: dict[str, float] = {}
    for key in ENTITY_ORDER:
        e = ENTITIES[key]
        mapping = suggest_mapping(key, headers)
        req = [f.key for f in e.fields if f.required]
        hits = sum(1 for r in req if r in mapping)
        score = sum(m["confidence"] for m in mapping.values()) + 1.5 * hits - 2.0 * (len(req) - hits)
        if any(k.replace(" ", "_") in fname for k in e.keywords):
            score += 4
        scores[key] = round(score, 2)
    best = max(ENTITY_ORDER, key=lambda k: (scores[k], -ENTITY_ORDER.index(k)))
    ranked = sorted(scores.values(), reverse=True)
    margin = ranked[0] - (ranked[1] if len(ranked) > 1 else 0)
    confidence = 0.95 if margin >= 4 else 0.75 if margin >= 2 else 0.5
    return best, confidence, scores


def validate_mapping(entity_key: str, mapping: dict[str, str], headers: list[str],
                     rows: list[list[str]]) -> dict:
    """mapping: {field_key: column}. Returns validation details for the UI."""
    e = ENTITIES[entity_key]
    problems, warnings = [], []
    col_idx = {h: i for i, h in enumerate(headers)}
    for fkey, col in mapping.items():
        if e.field(fkey) is None:
            problems.append(f"Unknown field '{fkey}' for {e.label.lower()}.")
        elif col not in col_idx:
            problems.append(f"Column '{col}' is not in this file.")
    missing = [f.label for f in e.fields if f.required and f.key not in mapping]
    for group in e.one_of:
        if not any(k in mapping for k in group):
            missing.append(" or ".join(e.field(k).label for k in group))
    if entity_key == "tenants" and not any(k in mapping for k in ("full_name", "first_name", "email", "phone")):
        warnings.append("No name, email or phone mapped — tenants can only be matched by ID.")
    stats = {"rows": len(rows), "blank_ids": 0, "duplicate_ids": 0, "bad_money": 0, "bad_dates": 0}
    id_col = mapping.get(e.id_field)
    if id_col in col_idx:
        seen = set()
        for r in rows:
            v = nz.ident(r[col_idx[id_col]])
            if v is None:
                stats["blank_ids"] += 1
            elif v in seen:
                stats["duplicate_ids"] += 1
            else:
                seen.add(v)
    for f in e.fields:
        col = mapping.get(f.key)
        if col not in col_idx:
            continue
        i = col_idx[col]
        if f.kind == "money":
            for r in rows:
                try:
                    nz.parse_money(r[i])
                except nz.MoneyError:
                    stats["bad_money"] += 1
        elif f.kind == "date":
            stats["bad_dates"] += sum(1 for r in rows if not nz.blank(r[i]) and nz.parse_date(r[i]) is None)
    if stats["blank_ids"]:
        warnings.append(f"{stats['blank_ids']} row(s) have a blank {e.field(e.id_field).label} and will be skipped."
                        if e.field(e.id_field).required else
                        f"{stats['blank_ids']} row(s) have no {e.field(e.id_field).label}; row numbers are used instead.")
    if stats["duplicate_ids"]:
        warnings.append(f"{stats['duplicate_ids']} duplicate ID(s) — flagged as a data-quality finding.")
    if stats["bad_money"]:
        warnings.append(f"{stats['bad_money']} amount(s) couldn't be read as money.")
    if stats["bad_dates"]:
        warnings.append(f"{stats['bad_dates']} date(s) couldn't be read.")
    status = "ready" if not problems and not missing else "needs_mapping"
    return {"status": status, "missing": missing, "problems": problems, "warnings": warnings, "stats": stats}
